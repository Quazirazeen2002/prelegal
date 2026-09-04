import io

import docx
import pytest
from fpdf import FPDF

from app.document_parsing import (
    EmptyDocumentError,
    FileTooLargeError,
    UnsupportedFileTypeError,
    extract_text,
    file_extension,
)


def _make_pdf_bytes(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text=text)
    return bytes(pdf.output())


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_file_extension() -> None:
    assert file_extension("contract.PDF") == "pdf"
    assert file_extension("no-extension") == ""


def test_extract_text_from_txt() -> None:
    assert extract_text("notes.txt", b"Hello, this is plain text.") == "Hello, this is plain text."


def test_extract_text_from_docx() -> None:
    content = _make_docx_bytes(["This Agreement is between Acme and Globex.", "Section 2. Term."])
    text = extract_text("agreement.docx", content)
    assert "This Agreement is between Acme and Globex." in text
    assert "Section 2. Term." in text


def test_extract_text_from_pdf() -> None:
    content = _make_pdf_bytes("This is a test contract.")
    text = extract_text("contract.pdf", content)
    assert "This is a test contract." in text


def test_unsupported_file_type_raises() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        extract_text("image.png", b"not a real file")


def test_file_too_large_raises() -> None:
    with pytest.raises(FileTooLargeError):
        extract_text("big.txt", b"x" * (101 * 1024 * 1024))


def test_empty_document_raises() -> None:
    with pytest.raises(EmptyDocumentError):
        extract_text("empty.txt", b"   \n\n  ")
